import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

df = pd.read_csv("data/netflix_cleaned.csv").drop(columns=["cast", "description"])
cols = ["show_id","type","title","director","country","date_added","release_year","year_added","rating","audience","duration_min","seasons","primary_country","primary_genre","decade"]
df = df[cols]
wb = Workbook(); F = "Arial"
hdr_fill = PatternFill("solid", fgColor="E50914"); hf = Font(name=F, bold=True, color="FFFFFF"); bf = Font(name=F)

# --- Data sheet
ws = wb.active; ws.title = "Cleaned_Data"
ws.append(cols)
for r in df.itertuples(index=False): ws.append([None if pd.isna(v) else v for v in r])
for c in ws[1]: c.font, c.fill = hf, hdr_fill
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
for i, w in enumerate([8,9,40,25,25,12,12,11,10,11,13,9,18,26,9], 1): ws.column_dimensions[L(i)].width = w
N = len(df) + 1
rng = lambda col: f"Cleaned_Data!${col}$2:${col}${N}"
T, Y, R, P, G, Au, Rt = rng("B"), rng("H"), rng("G"), rng("M"), rng("N"), rng("J"), rng("I")

# --- Dashboard sheet (all formulas)
d = wb.create_sheet("Summary", 0)
d["A1"] = "Netflix Content Dashboard — Summary (live formulas on Cleaned_Data)"; d["A1"].font = Font(name=F, bold=True, size=14, color="E50914")
kpis = [("Total Titles", f"=COUNTA({rng('A')})"), ("Movies", f'=COUNTIF({T},"Movie")'), ("TV Shows", f'=COUNTIF({T},"TV Show")'),
        ("% Movies", "=B4/B3"), ("Countries (primary)", None), ("Avg Movie Length (min)", f"=AVERAGE({rng('K')})")]
d["A2"] = "KPIs"; d["A2"].font = Font(name=F, bold=True)
for i, (k, f) in enumerate(kpis):
    if f: d.cell(3+i, 1, k); d.cell(3+i, 2, f)
d["A7"] = "Primary countries listed"; d["B7"] = "=SUMPRODUCT((" + P + '<>"Unknown")/COUNTIF(' + P + "," + P + '&""))'
d["B6"].number_format = "0.0%"; d["B8"].number_format = "0.0"
d["B6"] = "=B4/B3"; d["B6"].number_format = "0.0%"; d["A6"] = "% Movies"
d["A5"] = "TV Shows"; d["B5"] = f'=COUNTIF({T},"TV Show")'
d["A4"] = "Movies"; d["B4"] = f'=COUNTIF({T},"Movie")'; d["B6"] = "=B4/B3"
d["A8"] = "Avg Movie Length (min)"; d["B8"] = f"=AVERAGE({rng('K')})"

# Table: titles added per year
r0 = 11; d.cell(r0, 1, "Year Added"); d.cell(r0, 2, "Movies"); d.cell(r0, 3, "TV Shows"); d.cell(r0, 4, "Total"); d.cell(r0, 5, "YoY %")
years = list(range(2015, 2021))
for i, yr in enumerate(years):
    r = r0 + 1 + i; d.cell(r, 1, yr)
    d.cell(r, 2, f'=COUNTIFS({Y},A{r},{T},"Movie")'); d.cell(r, 3, f'=COUNTIFS({Y},A{r},{T},"TV Show")'); d.cell(r, 4, f"=B{r}+C{r}")
    if i: d.cell(r, 5, f"=IF(D{r-1}=0,0,D{r}/D{r-1}-1)").number_format = "0.0%"
d.cell(r0+8, 1, "Note: 2021 excluded (dataset ends Jan 2021, partial year).").font = Font(name=F, italic=True)

# Table: top countries / genres / ratings (labels from analysis, counts via formula)
cs = pd.read_csv("powerbi/title_country.csv").country.value_counts().head(10).index.tolist()
gs = pd.read_csv("powerbi/title_genre.csv").genre.value_counts().head(10).index.tolist()
rs = df.rating.value_counts().head(8).index.tolist()
tc = pd.read_csv("powerbi/title_country.csv"); tg = pd.read_csv("powerbi/title_genre.csv")
# helper sheets holding the bridge tables so COUNTIF formulas stay live
for nm, t_, col in [("Bridge_Country", tc, "country"), ("Bridge_Genre", tg, "genre")]:
    w = wb.create_sheet(nm); w.append(["show_id", col])
    for r in t_.itertuples(index=False): w.append(list(r))
    for c in w[1]: c.font, c.fill = hf, hdr_fill
def block(col0, title, items, fn):
    d.cell(r0+11, col0, title).font = Font(name=F, bold=True); d.cell(r0+11, col0+1, "Titles").font = Font(name=F, bold=True)
    for i, it in enumerate(items):
        r = r0 + 12 + i; d.cell(r, col0, it); d.cell(r, col0+1, fn(L(col0), r))
block(1, "Top Country", cs, lambda c, r: f"=COUNTIF(Bridge_Country!$B$2:$B${len(tc)+1},{c}{r})")
block(4, "Top Genre", gs, lambda c, r: f"=COUNTIF(Bridge_Genre!$B$2:$B${len(tg)+1},{c}{r})")
block(7, "Rating", rs, lambda c, r: f"=COUNTIF({Rt},{c}{r})")

for row in d.iter_rows():
    for c in row:
        if c.font == Font() or (c.font and c.font.name != F): c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, size=c.font.size if c.font.size else 11, color=c.font.color)
for c in [d.cell(r0, k) for k in range(1, 6)]: c.font, c.fill = hf, hdr_fill
for k in range(1, 9): d.column_dimensions[L(k)].width = 24
for s_ in (d,): s_.sheet_view.showGridLines = False

# charts
def bar(title, cat, vals, anchor, w=15, h=7.5):
    ch = BarChart(); ch.type = "col"; ch.title = title; ch.add_data(vals, titles_from_data=True); ch.set_categories(cat); ch.width, ch.height = w, h
    d.add_chart(ch, anchor)
bar("Titles Added per Year", Reference(d, min_col=1, min_row=r0+1, max_row=r0+6), Reference(d, min_col=2, max_col=3, min_row=r0, max_row=r0+6), "G2")
bar("Top 10 Countries", Reference(d, min_col=1, min_row=r0+12, max_row=r0+21), Reference(d, min_col=2, min_row=r0+11, max_row=r0+21), "A36")
bar("Top 10 Genres", Reference(d, min_col=4, min_row=r0+12, max_row=r0+21), Reference(d, min_col=5, min_row=r0+11, max_row=r0+21), "G36")
for nm in ("Cleaned_Data", "Bridge_Country", "Bridge_Genre"):
    for c in wb[nm][1]: c.font = hf

# --- Cleaning log
lg = wb.create_sheet("Cleaning_Log"); import json; rep = json.load(open("data/cleaning_report.json"))
rows = [("Step", "Action", "Result"),
 ("1", "Loaded raw Netflix titles CSV", f"{rep['raw_rows']:,} rows, 12 columns"),
 ("2", "Checked exact & title/type/year duplicates", f"{rep['exact_duplicates']} exact, {rep['title_dupes']} near-duplicates found (none to remove)"),
 ("3", "director / cast / country blanks -> 'Unknown'", f"{rep['missing_before']['director']:,} / {rep['missing_before']['cast']} / {rep['missing_before']['country']} filled"),
 ("4", "rating blanks -> 'Not Rated'", f"{rep['missing_before']['rating']} filled"),
 ("5", "date_added text ('August 14, 2020') -> DATE; year_added, month_added derived", f"{rep['date_added_missing']} blank dates kept as NULL (not imputed)"),
 ("6", "duration text split: '93 min' -> duration_min, '4 Seasons' -> seasons", "Two numeric columns, unit-consistent"),
 ("7", "Whitespace trimmed; multi-valued genre/country normalised into bridge tables", "title_genre, title_country"),
 ("8", "Derived: audience group, decade, primary_country, primary_genre", "Used for slicers")]
for r in rows: lg.append(r)
for c in lg[1]: c.font, c.fill = hf, hdr_fill
for r in lg.iter_rows(min_row=2):
    for c in r: c.font = bf
for k, w in zip("ABC", (7, 75, 65)): lg.column_dimensions[k].width = w
wb.save("Netflix_Analysis.xlsx")
